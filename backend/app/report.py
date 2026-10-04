"""Downloadable area report (D55): a PDF summary for officials and an Excel workbook, for an area or one street.

`content()` builds every number and table ONCE, from the same records and functions the app uses (the KPI formulas of
web/src/lib/derive.ts `kpis`, the app's plain labels from web/src/lib/labels.ts, the lighting priority of D54, Under
the Hood's routing and photo dates, Trust's Gate 1 wording and model-card numbers). `pdf()` and `xlsx()` only render
that content, so the two files carry the same columns and values. tests/test_report.py checks the numbers against the
API for Ward 29 and one small area.

No Google content is embedded: the map is drawn from our own data (OpenStreetMap roads from the app's copy, building
outlines, dark stretches) with "© OpenStreetMap contributors"; each row links to a normal Google Maps URL instead of an
image. Pure-Python libraries only (fpdf2, openpyxl, fontTools via fpdf2; uharfbuzz for Tamil text shaping).
"""
import datetime as _dt
import io
import math
import os

from . import camonly, lighting, mapdata, osmref, projection
from .hood import hood as hood_view
from .imagery import imagery as imagery_view

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "report_fonts")
LAMP_KEY = "lamp_head"
FOOTER = ("Registers are synthetic demo data. Prototype — imagery © Google (not reproduced here). "
          "Map data © OpenStreetMap contributors (ODbL).")


# ------------------------------------------------------------------------------------- the app's formatting and labels
def fmt(n):
    """Intl.NumberFormat('en-IN') for integers: 1,420 · 1,23,456"""
    if n is None:
        return "—"
    n = int(round(n))
    s, neg = str(abs(n)), n < 0
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        s = ",".join(groups + [tail])
    return ("-" if neg else "") + s


PLURALS = {"stretch": "stretches", "possible dark stretch": "possible dark stretches", "business": "businesses",
           "shop or business": "shops & businesses"}


def noun(n, one, many=None):
    return one if n == 1 else many or PLURALS.get(one) or one + "s"


def plural(n, one, many=None):
    return f"{fmt(n)} {noun(n, one, many)}"


def usd(v, digits=4):
    """lib/utils.ts usd()"""
    if v is None:
        return "—"
    step = 10 ** -digits
    if 0 < v < step / 2:
        return f"< ${step:.{digits}f}"
    return f"${v:.{digits}f}"


def usd_text(v):
    """Hood's usdText: 2 decimals from $1, 4 under 1 cent, else 3"""
    return "—" if v is None else usd(v, 2 if v >= 1 else 4 if v < 0.01 else 3)


MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def month_text(ym):
    if not ym:
        return "—"
    y, m = ym[:7].split("-")
    return f"{MONTHS[int(m) - 1]} {y}"


def day_text(iso):
    if not iso:
        return "date not recorded"
    d = _dt.date.fromisoformat(iso[:10])
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


MATCH = {"matched": "Matches the register", "discrepancy": "Differs from the register", "no_record": "Not in the register"}
MATCHED_USE_UNKNOWN = "Register entry exists — use not compared"
ASSET_REG = {"matched": "In the register", "discrepancy": "Differs from the register", "unrecorded_asset": "Not in the register",
             "unconfirmed_detection": "Seen in one photo only: needs a second look"}
DIFF = {"extra_floor": "more floors than recorded", "use_change": "used differently than recorded",
        "location_shift": "record pinned in the wrong place", "area_understated": "bigger than recorded",
        "missing_record": "no register record", "type_mismatch": "recorded as a different type"}
GAP_TYPE = {"poles present, no lamp detected": "poles, but no streetlight seen",
            "no pole or lamp detected": "no pole or streetlight seen"}
REVIEW = {"pending": "Waiting for review", "approved": "Approved", "rejected": "Rejected", "appealed": "Appealed"}
REASON_DIFF = {"missing_record": "Not in the register", "extra_floor": "Extra floor vs register",
               "use_change": "Use differs from register", "location_shift": "Register pin in the wrong place",
               "area_understated": "Bigger than recorded", "type_mismatch": "Recorded as a different type"}
PRIORITY_WORD = {"high": "High", "medium": "Medium", "low": "Low"}
ROAD_KIND_WORD = {"main": "main road", "connecting": "connecting road", "residential": "residential street",
                  "unknown": "not known"}


def pretty(s):
    return s.replace("_", " ") if s else "—"


def use_label(u):
    if not u:
        return "Use not known"
    return {"mixed": "Shop + home", "other": "Other", "under_construction": "Under construction"}.get(u, u[:1].upper() + u[1:])


def floors_text(n, status=None):
    if n is None:
        return "Floors not known"
    return f"{plural(n, 'floor')}{' (estimate)' if status == 'low_confidence' else ''}"


def match_label(s, use_known=True):
    if s == "matched" and not use_known:
        return MATCHED_USE_UNKNOWN
    return MATCH.get(s, pretty(s))


def review_reasons(reasons, finding=None):
    """lib/labels.ts reviewReasons"""
    import re
    finding = finding or {}
    diffs = [d for d in finding.get("discrepancies") or [] if d != "missing_record"]
    out = []

    def add(s):
        if s and s not in out:
            out.append(s)
    for r in reasons or []:
        if r == "high-severity discrepancy":
            if finding.get("match_status") == "no_record":
                add(REASON_DIFF["missing_record"])
            elif diffs:
                for d in diffs:
                    add(REASON_DIFF.get(d, pretty(d)))
            else:
                add("Differs from the register")
        elif re.match(r"^attribute discrepancy", r):
            a = [d for d in diffs if d in ("extra_floor", "use_change")]
            if a:
                for d in a:
                    add(REASON_DIFF[d])
            elif finding.get("match_status") == "no_record":
                add(REASON_DIFF["missing_record"])
            else:
                add("Differs from the register")
        elif r == "single-detection asset":
            add("Seen in one photo only")
        elif r == "building seen from one view only":
            add("Seen from one camera position only")
        elif re.match(r"^floor count low confidence", r):
            add("Floor count is an estimate (roof not visible)")
        elif re.match(r"^name read by VLM only", r):
            add("Shop name read by the AI model only")
        else:
            add(r[:1].upper() + re.sub(r"\s*[—–�-]\s*verify on imagery", "", r[1:]))
    return out


def maps_url(lat, lon):
    return f"https://www.google.com/maps/search/?api=1&query={lat:.6f},{lon:.6f}"


def loc(lat, lon):
    return f"{lat:.5f}, {lon:.5f}"


# ------------------------------------------------------------------------------------------------------- key numbers
KPI_DEFS = [  # web/src/lib/derive.ts KPI_DEFS: (key, label, singular label), the five main numbers first
    ("buildings_analysed", "Buildings checked", "Building checked"),
    ("unmatched_properties", "Not in register", None),
    ("buildings_with_discrepancy", "Differ from register", "Differs from register"),
    ("streetlight_gaps", "Possible dark stretches", "Possible dark stretch"),
    ("use_not_classified", "Use not known", None),
    ("streetlights", "Streetlights", "Streetlight"),
    ("poles", "Poles, no lamp seen", "Pole, no lamp seen"),
    ("named_businesses", "Shop names read clearly", "Shop name read clearly"),
    ("names_confirmed_by_google", "Also on Google Maps", None),
    ("sign_text_unverified", "Signs to double-check", "Sign to double-check"),
    ("waiting_for_review", "Waiting for review", None),
    ("unmapped_businesses", "Businesses with no analysed building", "Business with no analysed building"),
    ("streets_covered", "Streets", "Street"),
]
MAIN_KPIS = 5


def kpis(bundle, street=None):
    """derive.ts kpis(): the same formulas, optionally for one street"""
    on = lambda s: not street or s == street
    B = [b for b in bundle["buildings"] if on(b.get("street"))]
    A = [a for a in bundle["assets"] if on(a.get("street"))]
    nm = lambda b: ((b.get("attributes") or {}).get("name") or {})
    use = lambda b: ((b.get("attributes") or {}).get("use") or {}).get("value")
    R = [q for q in bundle["review_queue"] if on(q.get("street"))]
    return {
        "streets_covered": len({b.get("street") for b in B}),
        "buildings_analysed": len(B),
        "unmatched_properties": sum(b.get("match_status") == "no_record" for b in B),
        "buildings_with_discrepancy": sum(b.get("match_status") == "discrepancy" for b in B),
        "use_not_classified": sum(use(b) is None for b in B),
        "streetlights": sum(a.get("type") == "streetlight" for a in A),
        "poles": sum(a.get("type") == "pole" for a in A),
        "assets_triangulated": sum(a.get("method") == "triangulated" for a in A),
        "assets": len(A),
        "streetlight_gaps": sum(on(g.get("street")) for g in bundle["streetlight_gaps"]),
        "named_businesses": sum(nm(b).get("quality") == "good" for b in B),
        "sign_text_unverified": sum(nm(b).get("quality") in ("tamil_unverified", "fragment") for b in B),
        "names_confirmed_by_google": sum(bool(nm(b).get("google_confirmed")) for b in B),
        "low_confidence_observations": len(R),
        "waiting_for_review": sum(q.get("status") == "pending" for q in R),
        "unmapped_businesses": sum(on(u.get("street")) for u in bundle["unmapped_businesses"]),
    }


# ------------------------------------------------------------------------------------------------------------ tables
BUILDING_COLS = ["Building ID", "Street", "Approx. location (lat, lon)", "Use seen", "Floors seen", "Floor-count confidence",
                 "OSM building:levels (cross-check)", "Sign text", "In OpenStreetMap (shops)", "Confidence", "Finding",
                 "Register record (SYNTHETIC)", "Review status", "Google Maps"]
ASSET_COLS = ["Asset ID", "What", "Street", "Approx. location (lat, lon)", "Position", "Finding",
              "Register record (SYNTHETIC)", "Review status", "Google Maps"]
STRETCH_COLS = ["#", "Priority", "Why (plain words)", "Stretch ID", "Street", "Length (m, recorded)", "Road type (OSM)",
                "Shops & businesses ≤ 30 m", "Points (length + road + activity = total)", "Poles here", "Google Maps"]
REVIEW_COLS = ["Priority (1 = most urgent)", "Item", "Street", "Approx. location (lat, lon)", "Why it needs a look",
               "Status", "Google Maps"]


def _building_rows(bundle, street, mc=None, tags=None, shop_status=None):
    rank = {"no_record": 0, "discrepancy": 1}
    B = [b for b in bundle["buildings"] if b.get("match_status") in rank and (not street or b.get("street") == street)]
    B.sort(key=lambda b: (rank[b["match_status"]], b.get("street") or "", b["id"]))
    rows, links = [], []
    for b in B:
        at = b.get("attributes") or {}
        use, fl, nm = at.get("use") or {}, at.get("floors") or {}, at.get("name") or {}
        sv = (b.get("evidence") or {}).get("sign_view") or {}
        sign = nm.get("value") or sv.get("ocr_text") or "—"
        if nm.get("value") and nm.get("quality") and nm.get("quality") != "good":
            sign += {"fragment": " (partly readable)", "tamil_unverified": " (Tamil, not confirmed)"}.get(nm["quality"], "")
        reg = b.get("register") or {}
        conf = [{"measured": "floors measured", "low_confidence": "floors an estimate"}.get(fl.get("status"), "floors not measured")]
        if reg.get("match_confidence"):
            conf.append(f"register pairing {reg['match_confidence']}")
        finding = match_label(b["match_status"], bool(use.get("value")))
        diffs = [DIFF.get(d, pretty(d)) for d in b.get("discrepancies") or [] if d != "missing_record"]
        if diffs:
            finding += ": " + "; ".join(diffs)
        if reg.get("property_id"):
            parts = [str(reg["property_id"]), use_label(reg.get("record_use")) if reg.get("record_use") else "use not recorded",
                     floors_text(reg.get("record_floors")) if reg.get("record_floors") is not None else "floors not recorded"]
            if reg.get("record_area_m2") is not None:
                parts.append(f"{fmt(reg['record_area_m2'])} m²")
            if reg.get("record_dist_m") is not None:
                parts.append(f"pin {fmt(reg['record_dist_m'])} m away")
            record = " · ".join(parts)
        else:
            record = "No record"
        fc = osmref.floor_confidence(fl, mc)
        lv = osmref.building_levels(tags, b["id"])
        osm_lv = ("not loaded" if not lv["loaded"] else lv["osm_levels"] if lv["osm_levels"] is not None else "not tagged")
        shop = (shop_status or {}).get(b["id"], "—" if shop_status else "not loaded")
        rows.append([b["id"], b.get("street") or "—", loc(b["lat"], b["lon"]), use_label(use.get("value")),
                     floors_text(fl.get("value"), fl.get("status")), f"{fc['word']}: {fc['reason']}", osm_lv, sign, shop,
                     " · ".join(conf), finding, record,
                     REVIEW.get((b.get("review") or {}).get("status"), "Not queued"), "View in Google Maps"])
        links.append(maps_url(b["lat"], b["lon"]))
    return rows, links


def _asset_rows(bundle, street, everything=False):
    """poles and streetlights with a register finding (the PDF), or all of them with the finding blank when there is
    none (the Excel sheet, D56): the finding rows come first, identical in both"""
    keep = ("unrecorded_asset", "discrepancy")
    st = lambda a: (a.get("register") or {}).get("status")
    A = [a for a in bundle["assets"] if (everything or st(a) in keep) and (not street or a.get("street") == street)]
    A.sort(key=lambda a: (keep.index(st(a)) if st(a) in keep else len(keep), a.get("street") or "", a["id"]))
    rows, links = [], []
    for a in A:
        reg = a.get("register") or {}
        pos = ("pinpointed from " + plural(a.get("cameras_used") or 0, "camera position") if a.get("method") == "triangulated"
               else f"approximate (±{a.get('uncertainty_m'):g} m)" if a.get("uncertainty_m") is not None else "approximate")
        if reg.get("status") == "unconfirmed_detection":
            pos += " · seen in one photo only"
        finding = ASSET_REG.get(reg.get("status"), pretty(reg.get("status"))) if reg.get("status") in keep else ""
        if finding and reg.get("flags"):
            finding += ": " + "; ".join(DIFF.get(f, pretty(f)) for f in reg["flags"])
        record = " · ".join(x for x in (reg.get("asset_no"), pretty(reg.get("record_type")) if reg.get("record_type") else None) if x) \
            or "No record"
        rows.append([a["id"], "Streetlight" if a.get("type") == "streetlight" else "Pole, no lamp seen", a.get("street") or "—",
                     loc(a["lat"], a["lon"]), pos, finding, record,
                     REVIEW.get((a.get("review") or {}).get("status"), "Not queued"), "View in Google Maps"])
        links.append(maps_url(a["lat"], a["lon"]))
    return rows, links


def _stretch_rows(bundle, street, prio):
    gaps = {g["id"]: g for g in bundle["streetlight_gaps"] if not street or g.get("street") == street}
    if prio["available"]:
        order = [r["id"] for r in lighting.order([r for r in prio["rows"].values() if r["id"] in gaps])]
    else:
        order = [g["id"] for g in sorted(gaps.values(), key=lambda g: (-g["length_m"], g["id"]))]
    rows, links = [], []
    for i, gid in enumerate(order, 1):
        g, p = gaps[gid], prio["rows"].get(gid) if prio["available"] else None
        mid = ((g["start"][0] + g["end"][0]) / 2, (g["start"][1] + g["end"][1]) / 2)
        if p:
            pts = p["points"]
            road = ROAD_KIND_WORD[p["road_kind"]] + (f" ({p['road_class']})" if p["road_class"] else "")
            rows.append([i, PRIORITY_WORD[p["level"]], p["reason"], gid, g.get("street") or "—", fmt(g["length_m"]), road,
                         fmt(p["activity"]["total"]), f"{pts['length']} + {pts['road']} + {pts['activity']} = {p['score']}",
                         GAP_TYPE.get(g.get("gap_type"), g.get("gap_type") or "—"), "View in Google Maps"])
        else:
            rows.append([i, "not available", lighting.OFFLINE_NOTE, gid, g.get("street") or "—", fmt(g["length_m"]), "—", "—", "—",
                         GAP_TYPE.get(g.get("gap_type"), g.get("gap_type") or "—"), "View in Google Maps"])
        links.append(maps_url(*mid))
    return rows, links


def _review_rows(bundle, street):
    B = {b["id"]: b for b in bundle["buildings"]}
    A = {a["id"]: a for a in bundle["assets"]}
    Q = [q for q in bundle["review_queue"] if not street or q.get("street") == street]
    Q = sorted(Q, key=lambda q: (q.get("priority") or 9, q.get("street") or "", str(q.get("ref_id") or q.get("building_id") or "")))
    rows, links = [], []
    for q in Q:
        ref = q.get("ref_id") or q.get("building_id")
        if q.get("item_type") == "building":
            b = B.get(ref) or {}
            what = f"Building {ref} · {use_label(((b.get('attributes') or {}).get('use') or {}).get('value'))}"
            why = review_reasons(q.get("reasons"), {"match_status": b.get("match_status"), "discrepancies": q.get("discrepancies")
                                                    or b.get("discrepancies")})
        else:
            a = A.get(ref) or {}
            kind = a.get("type") or q.get("asset_cls")
            what = f"{'Streetlight' if kind == 'streetlight' else 'Pole'} {ref or ''}".strip()
            why = review_reasons(q.get("reasons"))
        rows.append([q.get("priority"), what, q.get("street") or "—", loc(q["lat"], q["lon"]), "; ".join(why) or "—",
                     REVIEW.get(q.get("status"), pretty(q.get("status"))), "View in Google Maps"])
        links.append(maps_url(q["lat"], q["lon"]))
    return rows, links


# ----------------------------------------------------------------------------------------------------- photo dates
def _street_photo_dates(F, names, street, today=None):
    """Hood's photo-date rule (imagery.py) for one street: the camera stops planned on it"""
    from .imagery import OLD_MONTHS, _today_m, months
    date = {p["pano_id"]: p.get("date") for p in F.get("panos") or [] if p.get("pano_id")}
    used = [date.get(e["pano_id"]) for e in F.get("plan") or [] if e.get("camera_lat") is not None
            and names.get(e.get("street"), e.get("street")) == street and date.get(e.get("pano_id"))]
    cutoff = _today_m(today) - OLD_MONTHS
    return {"camera_stops_dated": len(used), "oldest": min(used) if used else None, "newest": max(used) if used else None,
            "older_than_cutoff": sum(1 for d in used if months(d) is not None and months(d) < cutoff)}


def _street_names(areas_dir, slug):
    import json
    try:
        with open(os.path.join(areas_dir, slug, "street_names.json"), encoding="utf-8") as f:
            v = json.load(f)
        return v if isinstance(v, dict) else {}
    except (OSError, ValueError):
        return {}


# ------------------------------------------------------------------------------------------------------------ map
def _map(store, bundle, street, prio):
    """what the PDF map draws, in lon/lat: roads (OpenStreetMap copy; analysed streets when it is not available),
    the area outline, building outlines by finding, dark stretches by priority (numbered as in the table)"""
    on = lambda s: not street or s == street
    pts = []
    for b in bundle["buildings"]:
        if on(b.get("street")):
            pts += [(lo, la) for la, lo in ((b.get("footprint") or {}).get("polygon_latlon") or [[b["lat"], b["lon"]]])]
    for g in bundle["streetlight_gaps"]:
        if on(g.get("street")):
            pts += [(g["start"][1], g["start"][0]), (g["end"][1], g["end"][0])]
    for s in bundle["streets"]:
        if street and s["name"] == street and s.get("geometry"):
            geo = s["geometry"]
            lines = geo["coordinates"] if geo["type"] == "MultiLineString" else [geo["coordinates"]]
            pts += [tuple(p) for ln in lines for p in ln]
    if not street or not pts:
        w, s_, e, n = bundle["bbox"]
    else:
        xs, ys = [p[0] for p in pts], [p[1] for p in pts]
        w, s_, e, n = min(xs), min(ys), max(xs), max(ys)
    pad_y, pad_x = 40 / 110540, 40 / (111320 * math.cos(math.radians((s_ + n) / 2)))
    box = (w - pad_x, s_ - pad_y, e + pad_x, n + pad_y)
    roads, road_source = None, "analysed streets only (the OpenStreetMap copy was not available)"
    if getattr(store, "pool", None) is not None:
        ways = mapdata.roads_in_box(box[1], box[0], box[3], box[2])
        if ways is not None:
            roads = [{"cls": w_["tags"].get("highway"), "line": [(p["lon"], p["lat"]) for p in w_["geometry"]]} for w_ in ways]
            c = mapdata.city_for_box(box[1], box[0], box[3], box[2])
            road_source = f"OpenStreetMap snapshot {day_text(c['osm_snapshot'])} (the app's copy)" if c else "OpenStreetMap"
    if roads is None:
        roads = []
        for s in bundle["streets"]:
            geo = s.get("geometry")
            if geo:
                lines = geo["coordinates"] if geo["type"] == "MultiLineString" else [geo["coordinates"]]
                roads += [{"cls": "residential", "line": [tuple(p) for p in ln]} for ln in lines]
    buildings = [{"status": b.get("match_status"), "ring": [(lo, la) for la, lo in b["footprint"]["polygon_latlon"]]}
                 for b in bundle["buildings"] if on(b.get("street")) and len((b.get("footprint") or {}).get("polygon_latlon") or []) >= 3]
    disp = bundle.get("gap_display") or {}
    rank = {}
    if prio["available"]:
        rank = {r["id"]: i for i, r in enumerate(lighting.order([r for r in prio["rows"].values()
                                                                if on(next((g.get("street") for g in bundle["streetlight_gaps"]
                                                                            if g["id"] == r["id"]), None))]), 1)}
    gaps = []
    for g in bundle["streetlight_gaps"]:
        if not on(g.get("street")):
            continue
        d = disp.get(g["id"]) or {}
        line = d.get("path") if d.get("mode") == "along_road" and d.get("path") else [g["start"][::-1], g["end"][::-1]]
        p = prio["rows"].get(g["id"]) if prio["available"] else None
        gaps.append({"id": g["id"], "level": p["level"] if p else None, "rank": rank.get(g["id"]), "line": [tuple(x) for x in line]})
    outline = []
    poly = bundle.get("polygon") or {}
    if not street and poly.get("type") in ("Polygon", "MultiPolygon"):
        rings = [poly["coordinates"][0]] if poly["type"] == "Polygon" else [p[0] for p in poly["coordinates"]]
        outline = [[tuple(x) for x in r] for r in rings]
    return {"box": box, "roads": roads, "road_source": road_source, "buildings": buildings, "gaps": gaps, "outline": outline}


# ------------------------------------------------------------------------------------- extras: OSM, charts, actions
SHOP_COLS = ["Result", "Our business (ID)", "What we saw", "Street", "OpenStreetMap point", "OSM name", "OSM kind",
             "Distance (m)", "Same name", "Google Maps"]


def _shop_rows(sh):
    """the comparison as one table: matched, seen by our camera only, in OpenStreetMap only"""
    if not sh.get("available"):
        return [], []
    rows, links = [], []
    for m in sh["matched"]:
        c, o = m["camera"], m["osm"]
        rows.append(["Matched (same place)", c["id"], c.get("name") or c.get("sign_text") or c["why"], c.get("street") or "—",
                     o["osm_id"], o.get("name") or "—", o.get("kind"), m["distance_m"], "yes" if m["same_name"] else "no",
                     "View in Google Maps"])
        links.append(maps_url(c["lat"], c["lon"]))
    for c in sh["camera_only"]:
        rows.append(["Seen by our camera, not in OpenStreetMap", c["id"], c.get("name") or c.get("sign_text") or c["why"],
                     c.get("street") or "—", "—", "—", "—", None, "—", "View in Google Maps"])
        links.append(maps_url(c["lat"], c["lon"]))
    for o in sh["osm_only"]:
        rows.append(["In OpenStreetMap, not seen by our camera", "—", "—", "—", o["osm_id"], o.get("name") or "—", o.get("kind"),
                     None, "—", "View in Google Maps"])
        links.append(maps_url(o["lat"], o["lon"]))
    return rows, links


def _projection(store, areas_dir, mc):
    try:
        return projection.project(getattr(store, "pool", None), areas_dir, mc)
    except Exception as e:                                     # noqa: BLE001 - a report never fails on the projection
        return {"available": False, "note": f"not available ({type(e).__name__})"}


def _count(xs):
    out = {}
    for x in xs:
        out[x] = out.get(x, 0) + 1
    return sorted(out.items(), key=lambda kv: (-kv[1], str(kv[0])))


def _charts(bundle, street, sh):
    """the app's Charts (derive.ts charts()): building use, floor distribution (measured only, with n), register match
    status, findings by street, plus the OpenStreetMap shops comparison"""
    on = lambda s_: not street or s_ == street
    B = [b for b in bundle["buildings"] if on(b.get("street"))]
    use = lambda b: ((b.get("attributes") or {}).get("use") or {}).get("value")
    fl = lambda b: (b.get("attributes") or {}).get("floors") or {}
    measured = [b for b in B if fl(b).get("status") == "measured" and fl(b).get("value") is not None]
    floors = sorted(_count([fl(b)["value"] for b in measured]), key=lambda kv: kv[0])
    by = {}
    for b in B:
        r = by.setdefault(b.get("street") or "—", [0, 0, 0])
        r[0] += b.get("match_status") == "no_record"
        r[1] += b.get("match_status") == "discrepancy"
        r[2] += 1
    streets = sorted(by.items(), key=lambda kv: (-(kv[1][0] + kv[1][1]), kv[0]))
    return {
        "use": [(use_label(u) if u else "Use not known", n) for u, n in _count([use(b) for b in B])],
        "floors": [(plural(v, "floor"), n) for v, n in floors], "floors_n": len(measured),
        "floors_left_out": len(B) - len(measured),
        "match": [(match_label(m_), n) for m_, n in _count([b.get("match_status") for b in B])],
        "by_street": [(st, r[0], r[1], r[2]) for st, r in streets],
        "shops": ([("Matched (same place)", sh["counts"]["matched"]), ("Seen by our camera only", sh["counts"]["camera_only"]),
                   ("In OpenStreetMap only", sh["counts"]["osm_only"])] if sh.get("available") else None),
    }


def _headline(k, ch, s_rows, prio_ok, sh, street):
    """three plain sentences: what matters most, computed from the numbers"""
    out = []
    top = next(((st, a + d) for st, a, d, _ in ch["by_street"] if a + d), None)
    if k["buildings_analysed"]:
        line = (f"{fmt(k['unmatched_properties'])} of the {plural(k['buildings_analysed'], 'building')} checked "
                f"{'is' if k['unmatched_properties'] == 1 else 'are'} not in the register and {fmt(k['buildings_with_discrepancy'])} "
                f"{'differs' if k['buildings_with_discrepancy'] == 1 else 'differ'} from it (synthetic register)")
        if top and not street and len(ch["by_street"]) > 1:
            line += f"; the most are on {top[0]} ({fmt(top[1])})"
        out.append(line + ".")
    else:
        out.append("No building with a map outline was analysed here.")
    n = len(s_rows)
    if n:
        first = s_rows[0]
        hi = sum(1 for r in s_rows if r[1] == "High")
        out.append(f"{plural(n, 'possible dark stretch')} {'was' if n == 1 else 'were'} found"
                   + (f", {fmt(hi)} of {'it' if n == 1 else 'them'} high priority" if prio_ok else "")
                   + f"; check {first[4]} first ({first[5]} m).")
    else:
        out.append("No possible dark stretch was found: a streetlight was seen at least every 60 m.")
    if sh.get("available") and sh["counts"]["camera"]:
        c = sh["counts"]
        out.append(f"Our camera found {plural(c['camera'], 'shop or business', 'shops and businesses')}; "
                   f"{fmt(c['camera_only'])} of them {'is' if c['camera_only'] == 1 else 'are'} not on OpenStreetMap"
                   + (f", and {fmt(c['osm_only'])} OpenStreetMap {noun(c['osm_only'], 'shop')} "
                      f"{'was' if c['osm_only'] == 1 else 'were'} not seen by the camera." if c["osm_only"] else "."))
    else:
        out.append(f"{plural(k['waiting_for_review'], 'item')} {'waits' if k['waiting_for_review'] == 1 else 'wait'} "
                   "for a person to check.")
    return out


def _actions(bundle, street, s_rows, sh, k, recall, r_rows):
    """the top ~10 things to do, with street names and one line of why each"""
    acts = []
    for r in [r for r in s_rows if r[1] == "High"][:3]:
        acts.append((f"Check the street lighting on {r[4]} ({r[5]} m)",
                     f"High-priority possible dark stretch: {r[2]}. No lamp was seen there; the detector finds about {recall} "
                     "of lamp heads, so look on site before ordering work."))
    if not acts and s_rows:
        r = s_rows[0]
        acts.append((f"Check the street lighting on {r[4]} ({r[5]} m)",
                     "The first possible dark stretch in the list: no lamp was seen within 60 m."))
    on = lambda s_: not street or s_ == street
    nr, df = {}, {}
    for b in bundle["buildings"]:
        if on(b.get("street")):
            if b.get("match_status") == "no_record":
                nr.setdefault(b.get("street") or "—", []).append(b)
            elif b.get("match_status") == "discrepancy":
                df.setdefault(b.get("street") or "—", []).append(b)
    for st, bs in sorted(nr.items(), key=lambda kv: (-len(kv[1]), kv[0]))[:3]:
        acts.append((f"Verify {plural(len(bs), 'building')} on {st} that {'is' if len(bs) == 1 else 'are'} not in the register",
                     "The camera saw them, but no register record lies near them (synthetic register: with a real register "
                     "these would be properties missing from it)."))
    for st, bs in sorted(df.items(), key=lambda kv: (-len(kv[1]), kv[0]))[:2]:
        kinds = _count([d for b in bs for d in (b.get("discrepancies") or []) if d != "missing_record"])
        acts.append((f"Check {plural(len(bs), 'register record')} on {st} that {'differs' if len(bs) == 1 else 'differ'} "
                     "from the street",
                     ("Most often: " + ", ".join(f"{DIFF.get(d, pretty(d))} ({n})" for d, n in kinds[:2]) + ".") if kinds else
                     "The record does not match what the photos show."))
    if sh.get("available") and sh["camera_only"]:
        named = [c for c in sh["camera_only"] if c.get("name")]
        st, n = _count([c.get("street") or "—" for c in (named or sh["camera_only"])])[0]
        acts.append((f"Add {plural(n, 'shop')} seen on {st} to OpenStreetMap",
                     ("Their signs were read clearly" if named else "The camera saw them")
                     + f", but OpenStreetMap has no shop point within {sh['match_m']} m (it is the open map anyone can edit)."))
    p1 = [r for r in r_rows if r[0] == 1 and r[5] == "Waiting for review"]
    if k["waiting_for_review"]:
        acts.append((f"Work through the review queue: {plural(k['waiting_for_review'], 'item')} waiting",
                     f"Start with the {fmt(len(p1))} most urgent (priority 1); each has its Street View evidence in the app."
                     if p1 else "Each has its Street View evidence in the app."))
    return [{"what": a, "why": w} for a, w in acts[:10]]


# --------------------------------------------------------------------------------------------------------- content
def content(store, bundle, F, model_card, areas_dir, street=None, today=None):
    if street and street not in {b.get("street") for b in bundle["buildings"]} | {s["name"] for s in bundle["streets"]} \
            | {g.get("street") for g in bundle["streetlight_gaps"]} | {a.get("street") for a in bundle["assets"]}:
        raise KeyError(street)
    mc = model_card or {}
    today = today or _dt.date.today()
    area_name = bundle["name"].replace("Unseen street: ", "")
    k = kpis(bundle, street)
    numbers = [{"key": key, "label": (one if k[key] == 1 and one else label), "value": k[key], "main": i < MAIN_KPIS}
               for i, (key, label, one) in enumerate(KPI_DEFS)]
    cam = 0 if street else camonly.count(bundle["slug"], areas_dir)          # D56: next to the count, never added to it
    numbers[0]["note"] = camonly.phrase(cam)
    prio = lighting.for_area(store, bundle)
    hood = hood_view(bundle, F, mc)
    im = imagery_view(bundle, F)["summary"] if not street else _street_photo_dates(F, _street_names(areas_dir, bundle["slug"]), street)
    md = mapdata.area_map_data(bundle, areas_dir)
    lamp = ((mc.get("detector") or {}).get("per_class") or {}).get(LAMP_KEY) or {}
    recall = f"{round(lamp['R'] * 100)}%" if lamp.get("R") is not None else "not measured"

    # data sources
    city = md.get("city")
    run = md.get("run") or {}
    sources = [
        ("Street View photos (Google)", (f"taken {month_text(im['oldest'])} – {month_text(im['newest'])}; "
                                         f"{im['older_than_cutoff']} of {im['camera_stops_dated']} camera positions are over 3 years old")
         if im.get("oldest") else "photo dates not recorded"),
        ("OpenStreetMap (roads, building outlines, shop points)",
         (f"this analysis used OpenStreetMap's snapshot of {day_text(run.get('osm_snapshot'))}" if run.get("kind") == "snapshot"
          else f"this analysis read OpenStreetMap's live servers{(' on ' + day_text(run['date'])) if run.get('date') else ' on its run day'}")
         + (f"; this report's map and road types use the app's snapshot of {day_text(city['osm_snapshot'])}" if city else "")),
        ("Property and asset registers", "SYNTHETIC demo data: made-up records that copy what the photos show, except planted "
                                         "mistakes. Not the city's records."),
        ("Accuracy figures", "the team's model card (measured on hand-labelled samples)"),
    ]
    if city and city.get("ms_release"):
        sources.insert(2, ("Microsoft building footprints", f"release {day_text(city['ms_release'])} (used only where OpenStreetMap has few outlines)"))

    # routing and cost (Under the Hood › Routing and cost)
    rt = hood.get("routing")
    cost = None
    if rt:
        sv, ai = rt["street_view"], rt["totals"]["usd"]
        total = sv["usd"] + ai if sv.get("usd") is not None and ai is not None else None
        line = None
        if total:
            cloud = sorted([(x["usd"], t["key"]) for t in rt["tasks"] for x in t["routes"] if x["route"] == "cloud" and x["usd"] is not None],
                           reverse=True)
            line = (f"This run cost about {usd_text(total)}: Street View photos {usd_text(sv['usd'])} ({fmt(sv['photos'])} at "
                    f"Google’s list price, {round(100 * sv['usd'] / total)}%) and cloud AI (Nova Lite) {usd_text(ai)} "
                    f"({100 * ai / total:.1f}%)." + (" Floor counting is the largest AI cost; it isn’t routed yet."
                                                     if cloud and cloud[0][1] == "floors" else ""))
        tasks = [[t["title"], f"{x['route']} · {x['model']}", fmt(x["n"]), "$0" if x["route"] == "local" else usd_text(x["usd"]),
                  "" if x["route"] == "local" else x["usd_status"]] for t in rt["tasks"] if t["input"] > 0 for x in t["routes"]]
        cost = {"line": line, "tasks": tasks, "as_run": (rt["totals"]["calls"], rt["totals"]["usd"], rt["totals"]["status"]),
                "no_router": (rt["all_cloud"]["calls"], rt["all_cloud"]["usd"]) if rt.get("all_cloud") else None}

    # Gate 1, worded as on Trust
    g1 = mc.get("gate1_position") or {}
    fair = None
    for key, v in ((g1.get("vs OSM front-wall centre") or {}).get(bundle["slug"]) or {}).items():
        if key.startswith("camera-derived") and v.get("n"):
            fair = v
    gate1 = {"title": f"Position accuracy — target ≤ {g1.get('target_m')} m (FarmwiseAI Gate 1)",
             "status": "Status: " + ("Not verified" if g1.get("status") == "not verified" else str(g1.get("status"))),
             "note": f"{g1.get('status_note')} No surveyed reference exists.",
             "row": (f"All camera-derived positions (the two camera methods together), vs the centre of the OSM front wall: "
                     f"n = {fmt(fair['n'])}, median {fair['median_m']} m, p90 {fair['p90_m']} m, {fair['within_3_5_m_pct']}% within "
                     f"{g1.get('target_m')} m.") if fair else "This area is not in Trust's Gate 1 table (it was analysed from the app)."}

    limits = [
        "The property and asset registers are SYNTHETIC. They copy what the photos show, except planted mistakes, so "
        "“not in the register” and “differs” test the comparison end to end; they do not say the city’s records are wrong.",
        f"Possible dark stretches mean no lamp was seen within 60 m, not that none exists: the lamp detector finds about {recall} of "
        f"lamp heads (n = {lamp.get('n', '—')}), and a photo can’t tell whether a lamp works. The priority ranks these candidates; "
        "it does not confirm them.",
        "The priority’s activity count includes buildings whose use is commercial or shop + home or that carry a shop name "
        f"read clearly ({plural(k['use_not_classified'], 'building')} here have no known use), OpenStreetMap shop / amenity / "
        "office points (which include some non-business places) and businesses read from signs.",
        (f"Street View photos are {month_text(im['oldest'])} – {month_text(im['newest'])}; {im['older_than_cutoff']} of "
         f"{im['camera_stops_dated']} camera positions are over 3 years old, so the street may have changed since.")
        if im.get("oldest") else "Photo dates are not recorded for this area.",
        "Gate 1 (building position within 3.5 m) is measured against OpenStreetMap’s outlines, not a surveyed reference, so it "
        "is not verified.",
        "Locations are approximate (map points, not survey points). The Google Maps links open the location only.",
        "OpenStreetMap is crowd-sourced (volunteers), not an official register: a shop or a floor count missing from it says "
        "nothing about the street. It is shown only as a cross-check and changes none of our results.",
    ]

    tags = osmref.load(bundle["slug"], areas_dir)
    shops = osmref.shops(bundle, tags, street)
    lv = osmref.levels(bundle, tags)
    proj = _projection(store, areas_dir, mc)
    b_rows, b_links = _building_rows(bundle, street, mc, tags, osmref.shop_status_by_id(shops))
    a_rows, a_links = _asset_rows(bundle, street)
    aa_rows, aa_links = _asset_rows(bundle, street, everything=True)
    s_rows, s_links = _stretch_rows(bundle, street, prio)
    r_rows, r_links = _review_rows(bundle, street)
    sh_rows, sh_links = _shop_rows(shops)
    charts = _charts(bundle, street, shops)
    pcounts = {lv_: sum(1 for r in s_rows if r[1] == PRIORITY_WORD[lv_]) for lv_ in PRIORITY_WORD}
    return {
        "area": area_name, "slug": bundle["slug"], "street": street, "generated": today.isoformat(),
        "title": f"{street} — {area_name}" if street else area_name,
        "scope": f"One street: {street}" if street else f"Whole area ({plural(k['streets_covered'], 'street')})",
        "sources": sources, "numbers": numbers, "kpis": k, "recall": recall, "camera_only_buildings": cam,
        "priority": {"available": prio["available"], "rule": prio["rule"], "note": prio["note"],
                     "counts": {lv: sum(1 for r in s_rows if r[1] == PRIORITY_WORD[lv]) for lv in PRIORITY_WORD}},
        "tables": {
            "buildings": {"title": "Buildings with findings", "sheet": "Buildings with findings", "columns": BUILDING_COLS,
                          "rows": b_rows, "links": b_links},
            "assets": {"title": "Poles and streetlights with a register finding", "sheet": "Assets", "columns": ASSET_COLS,
                       "rows": a_rows, "links": a_links},
            "assets_all": {"title": "All poles and streetlights (register finding blank when there is none)", "sheet": "Assets",
                           "columns": ASSET_COLS, "rows": aa_rows, "links": aa_links},
            "stretches": {"title": "Possible dark stretches, in priority order", "sheet": "Possible dark stretches",
                          "columns": STRETCH_COLS, "rows": s_rows, "links": s_links},
            "review": {"title": "Review queue", "sheet": "Review items", "columns": REVIEW_COLS, "rows": r_rows, "links": r_links},
            "shops": {"title": "Businesses: our camera vs OpenStreetMap", "sheet": "OSM shops", "columns": SHOP_COLS,
                      "rows": sh_rows, "links": sh_links},
        },
        "cost": cost, "cost_scope": "for the whole analysed area’s run" if street else None,
        "gate1": gate1, "limits": limits, "map": _map(store, bundle, street, prio),
        "shops": shops, "levels": lv, "projection": proj, "charts": charts, "floor_rule": osmref.FLOOR_RULE,
        "floor_check": osmref.floor_confidence({"value": 1, "status": "measured"}, mc)["check"],
        "city": (md.get("city") or {}).get("name") if md.get("city") else None,
        "headline": _headline(k, charts, s_rows, prio["available"], shops, street),
        "actions": _actions(bundle, street, s_rows, shops, k, recall, r_rows),
        "priority_counts": pcounts,
    }


# ----------------------------------------------------------------------------------------------------------- PDF
INK, INK2, INK3 = (27, 31, 42), (61, 67, 86), (91, 96, 112)
LINE = (214, 210, 200)
SODIUM = (156, 78, 7)
STATUS = {"no_record": (204, 31, 99), "discrepancy": (0, 132, 159), "matched": (142, 157, 196)}
PRIORITY_RGB = {"high": (42, 42, 143), "medium": (85, 96, 204), "low": (143, 153, 227)}   # daylight ramp (D54; re-stepped D57)
DARK = (27, 31, 42)
ROAD_W = {"motorway": 1.2, "trunk": 1.2, "primary": 1.0, "secondary": 0.9, "tertiary": 0.7}


def _hex(rgb):
    return "%02X%02X%02X" % rgb


SYMBOLS = {"≤": "<=", "≥": ">=", "→": "->", "↑": "^"}   # not in Anek Tamil; drawn from Arial when the system has it
SYSTEM_SYMBOL_FONT = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "arial.ttf")


def _font(pdf):
    for w, style in ((400, ""), (600, "B")):
        pdf.add_font("Anek", style, os.path.join(FONT_DIR, f"AnekTamil-latin-{w}.ttf"))
        pdf.add_font("AnekTa", style, os.path.join(FONT_DIR, f"AnekTamil-tamil-{w}.ttf"))
    fallbacks = ["AnekTa"]
    if os.path.isfile(SYSTEM_SYMBOL_FONT):
        for style in ("", "B"):
            pdf.add_font("Sym", style, SYSTEM_SYMBOL_FONT)
        fallbacks.append("Sym")
    else:
        pdf.replace_symbols = True
    pdf.set_fallback_fonts(fallbacks, exact_match=False)
    try:
        pdf.set_text_shaping(True)                     # Tamil vowel signs need shaping (uharfbuzz)
    except Exception:                                  # noqa: BLE001 - without uharfbuzz Latin text is unchanged
        pass


def _pdf_class():
    from fpdf import FPDF

    class Report(FPDF):
        doc_title = ""
        replace_symbols = False

        def normalize_text(self, text):
            if self.replace_symbols:
                for a, b in SYMBOLS.items():
                    text = text.replace(a, b)
            return super().normalize_text(text)

        def footer(self):
            self.set_y(-13)
            self.set_font("Anek", "", 8.5)
            self.set_text_color(*INK3)
            self.cell(0, 4, f"{self.doc_title} · page {self.page_no()} of {{nb}}", align="L", new_x="LMARGIN", new_y="NEXT")
            self.cell(0, 4, FOOTER, align="L")
    return Report


def _draw_map(pdf, m, x, y, w, h, small=False):
    w0, s0, e0, n0 = m["box"]
    kx = math.cos(math.radians((s0 + n0) / 2))
    gw, gh = (e0 - w0) * kx, (n0 - s0)
    sc = min(w / gw, h / gh)
    ox, oy = x + (w - gw * sc) / 2, y + (h - gh * sc) / 2
    P = lambda lon, lat: (ox + (lon - w0) * kx * sc, oy + (n0 - lat) * sc)
    pdf.set_fill_color(250, 248, 242)
    pdf.set_draw_color(*LINE)
    pdf.set_line_width(0.2)
    pdf.rect(x, y, w, h, style="DF")
    with pdf.rect_clip(x, y, w, h):
        pdf.set_draw_color(190, 186, 176)
        for r in m["roads"]:
            if len(r["line"]) >= 2:
                pdf.set_line_width(ROAD_W.get((r["cls"] or "").replace("_link", ""), 0.45))
                pdf.polyline([P(*p) for p in r["line"]])
        for ring in m["outline"]:
            pdf.set_draw_color(*INK3)
            pdf.set_line_width(0.35)
            with pdf.local_context(dash_pattern={"dash": 1.2, "gap": 0.8}):
                pdf.polyline([P(*p) for p in ring])
        pdf.set_line_width(0.12)
        for status in ("matched", "discrepancy", "no_record"):                 # findings drawn on top
            col = STATUS.get(status)
            pdf.set_fill_color(*col)
            pdf.set_draw_color(*[max(0, c - 40) for c in col])
            for b in m["buildings"]:
                if b["status"] == status:
                    pdf.polygon([P(*p) for p in b["ring"]], style="DF")
        for lv in ("low", "medium", "high", None):
            for g in m["gaps"]:
                if g["level"] != lv or len(g["line"]) < 2:
                    continue
                pts = [P(*p) for p in g["line"]]
                lw = {"high": 2.4, "medium": 2.0, "low": 1.7}.get(lv, 1.7)
                pdf.set_draw_color(255, 255, 255)                                  # D57: white casing, as on the map
                pdf.set_line_width(lw + 0.7)
                pdf.polyline(pts)
                pdf.set_draw_color(*(PRIORITY_RGB.get(lv) or (150, 150, 150)))
                pdf.set_line_width({"high": 2.4, "medium": 2.0, "low": 1.7}.get(lv, 1.7))
                pdf.polyline(pts)
                pdf.set_draw_color(*DARK)
                pdf.set_line_width(0.9)
                pdf.polyline(pts)
        for g in m["gaps"]:                                                        # rank numbers = the table's "#"
            if g["rank"] is None:
                continue
            cx, cy = P(*g["line"][len(g["line"]) // 2])
            pdf.set_fill_color(255, 255, 255)
            pdf.set_draw_color(*(PRIORITY_RGB.get(g["level"]) or INK))
            pdf.set_line_width(0.4)
            r_ = 2.1 if small else 2.8
            pdf.circle(x=cx, y=cy, radius=r_, style="DF")
            pdf.set_font("Anek", "B", 6.5 if small else 8.5)
            pdf.set_text_color(*INK)
            pdf.set_xy(cx - r_, cy - r_ * 0.76)
            pdf.cell(2 * r_, 1.52 * r_, str(g["rank"]), align="C")
    # scale bar + attribution
    m_per_mm = 111320 / sc
    nice = next(v for v in (10, 20, 50, 100, 200, 250, 500, 1000, 2000) if v / m_per_mm >= 15)
    L = nice / m_per_mm
    pdf.set_draw_color(*INK)
    pdf.set_line_width(0.4)
    pdf.line(x + 4, y + h - 5, x + 4 + L, y + h - 5)
    pdf.set_font("Anek", "", 8 if small else 9.5)
    pdf.set_text_color(*INK2)
    pdf.set_xy(x + 4, y + h - 10)
    pdf.cell(L + 10, 4, f"{fmt(nice)} m")
    pdf.set_xy(x + w - 72, y + h - 6)
    pdf.cell(70, 4.5, "© OpenStreetMap contributors", align="R")
    pdf.set_xy(x + w - 10, y + 2)
    pdf.set_font("Anek", "B", 9 if small else 10.5)
    pdf.cell(8, 4, "N ↑", align="C")


def _table(pdf, t, widths, size=9, link_col=True):
    from fpdf.fonts import FontFace
    pdf.set_font("Anek", "", size)
    pdf.set_text_color(*INK)
    pdf.set_draw_color(*LINE)
    pdf.set_line_width(0.2)
    pdf.set_fill_color(255, 255, 255)
    head = FontFace(emphasis="BOLD", color=INK, fill_color=(236, 232, 222))
    if not t["rows"]:
        _text(pdf, "None.", color=INK3)
        return
    with pdf.table(col_widths=widths, headings_style=head, line_height=LH(size), padding=1.2,
                   borders_layout="HORIZONTAL_LINES", first_row_as_headings=True, text_align="LEFT",
                   cell_fill_mode="NONE") as tb:
        r = tb.row()
        for c in t["columns"]:
            r.cell(c)
        for row, link in zip(t["rows"], t["links"]):
            r = tb.row()
            for i, v in enumerate(row):
                if link_col and i == len(row) - 1 and link:
                    r.cell(str(v), link=link, style=FontFace(color=SODIUM, emphasis="UNDERLINE"))
                else:
                    r.cell("" if v is None else fmt(v) if isinstance(v, int) and not isinstance(v, bool) else str(v))


def pdf(c):
    """extras 5: every page landscape A4, readable in two minutes. Main part (about five pages): at a glance · charts ·
    map · what to do next · method & limits. Then the appendix tables (key columns; every column in the Excel file).
    Body text 11 pt, tables 9 pt, headline numbers 30 pt."""
    Report = _pdf_class()
    doc = Report(orientation="L", unit="mm", format="A4")
    doc.doc_title = f"GEO-CASCADIA report · {c['title']}"
    doc.set_title(doc.doc_title)
    doc.set_author("GEO-CASCADIA (prototype)")
    doc.set_creator("GEO-CASCADIA API")
    _font(doc)
    doc.set_auto_page_break(True, margin=M + 6)
    doc.set_margins(M, M, M)
    doc.alias_nb_pages()
    W = doc.w - 2 * M
    doc.set_auto_page_break(False)                 # the main pages are laid out by hand to fit one page each
    _page_glance(doc, c, W)
    _page_charts(doc, c, W)
    short = bool(c["street"]) and len(c["actions"]) <= 6        # a street report: the actions sit beside the map
    _page_map(doc, c, W, actions=short)
    if not short:
        _page_actions(doc, c, W)
    _page_method(doc, c, W)
    doc.set_auto_page_break(True, margin=M + 6)
    _appendix(doc, c, W)
    return bytes(doc.output())


M = 14                                        # page margin, mm (all four sides)
BODY, SMALL, TABLE = 11, 9.5, 9               # pt
LH = lambda pt: pt * 0.3528 * 1.38            # line height in mm for a font size in pt
SOFT = (246, 243, 236)                        # card / panel fill
BAR = (85, 96, 204)                           # one hue for single-series bars (the Daylight indigo, D57)


def _kicker(doc, text):
    doc.set_font("Anek", "B", 10)
    doc.set_text_color(*SODIUM)
    doc.cell(0, 5, text.upper(), new_x="LMARGIN", new_y="NEXT")


def _title(doc, text, size=20, after=2):
    doc.set_font("Anek", "B", size)
    doc.set_text_color(*INK)
    doc.multi_cell(0, LH(size), text, new_x="LMARGIN", new_y="NEXT", align="L")
    doc.ln(after)


def _text(doc, text, size=BODY, color=INK2, bold=False, w=0, after=1.5):
    doc.set_font("Anek", "B" if bold else "", size)
    doc.set_text_color(*color)
    doc.multi_cell(w, LH(size), text, new_x="LMARGIN" if not w else "LEFT", new_y="NEXT", align="L")
    doc.ln(after)


def _in_column(doc, x, w):
    """write inside a column: left margin moves to x, the width is w"""
    doc.set_left_margin(x)
    doc.set_right_margin(doc.w - x - w)
    doc.set_x(x)


def _reset_columns(doc):
    doc.set_left_margin(M)
    doc.set_right_margin(M)


def _card(doc, x, y, w, h, value, label, note=None, color=INK):
    doc.set_fill_color(*SOFT)
    doc.set_draw_color(*LINE)
    doc.set_line_width(0.2)
    doc.rect(x, y, w, h, style="DF", round_corners=True, corner_radius=2)
    doc.set_xy(x + 4, y + 3)
    doc.set_font("Anek", "B", 30)
    doc.set_text_color(*color)
    doc.cell(w - 8, 12, value)
    doc.set_xy(x + 4, y + 16)
    doc.set_font("Anek", "B", 11.5)
    doc.set_text_color(*INK)
    doc.multi_cell(w - 8, LH(11.5), label, align="L")
    if note:
        doc.set_x(x + 4)
        doc.set_font("Anek", "", 9.5)
        doc.set_text_color(*INK2)
        doc.multi_cell(w - 8, LH(9.5), note, align="L")


def _page_glance(doc, c, W):
    doc.add_page()
    _kicker(doc, "GEO-CASCADIA · street survey report · at a glance")
    _title(doc, c["title"], 22, after=0.5)
    _text(doc, f"{c['scope']} · report made {day_text(c['generated'])}", size=BODY, color=INK2, after=2)
    map_w = 96
    left_w = W - map_w - 8
    y0 = doc.get_y()
    # the synthetic banner
    doc.set_fill_color(253, 236, 242)
    doc.set_draw_color(*STATUS["no_record"])
    doc.set_line_width(0.3)
    doc.set_font("Anek", "B", BODY)
    doc.set_text_color(*INK)
    doc.set_xy(M, y0)
    doc.multi_cell(left_w, LH(BODY), "The property and asset registers in this report are SYNTHETIC demo data (made up), "
                                      "not the city's records.", border=1, fill=True, padding=1.6, new_x="LMARGIN", new_y="NEXT")
    # number cards, 3 x 2
    k = c["kpis"]
    pc = c["priority_counts"]
    nums = {n["key"]: n for n in c["numbers"]}
    sh = c["shops"]
    cards = [
        (fmt(k["buildings_analysed"]), nums["buildings_analysed"]["label"], nums["buildings_analysed"].get("note") or None, INK),
        (fmt(k["unmatched_properties"]), "Not in the register", "synthetic register", STATUS["no_record"]),
        (fmt(k["buildings_with_discrepancy"]), "Differ from the register" if k["buildings_with_discrepancy"] != 1 else
         "Differs from the register", "synthetic register", STATUS["discrepancy"]),
        (fmt(k["streetlight_gaps"]), nums["streetlight_gaps"]["label"],
         (f"High {pc['high']} · Medium {pc['medium']} · Low {pc['low']}" if c["priority"]["available"] else "priority not available")
         + " · no lamp seen within 60 m", INK),
        (fmt(k["low_confidence_observations"]), "Review items" if k["low_confidence_observations"] != 1 else "Review item",
         f"{fmt(k['waiting_for_review'])} waiting for a person", INK),
        ((fmt(sh["counts"]["camera_only"]), "Businesses not on OpenStreetMap",
          f"of the {fmt(sh['counts']['camera'])} our camera found", INK) if sh.get("available") else
         (fmt(k["named_businesses"]), nums["named_businesses"]["label"], None, INK)),
    ]
    gap = 4
    cw = (left_w - 2 * gap) / 3
    ch = 38
    yc = doc.get_y() + 3
    for i, (v, lab, note, col) in enumerate(cards):
        _card(doc, M + (i % 3) * (cw + gap), yc + (i // 3) * (ch + gap), cw, ch, v, lab, note, col)
    doc.set_xy(M, yc + 2 * ch + gap + 6)
    _in_column(doc, M, left_w)
    doc.set_font("Anek", "B", 13)
    doc.set_text_color(*INK)
    doc.cell(0, 6, "What matters most", new_x="LMARGIN", new_y="NEXT")
    doc.ln(1)
    for i, s_ in enumerate(c["headline"], 1):
        doc.set_font("Anek", "B", BODY + 0.5)
        doc.set_text_color(*SODIUM)
        doc.cell(6, LH(BODY + 0.5), f"{i}.")
        doc.set_font("Anek", "", BODY + 0.5)
        doc.set_text_color(*INK)
        doc.multi_cell(left_w - 6, LH(BODY + 0.5), s_, new_x="LMARGIN", new_y="NEXT", align="L")
        doc.ln(1.2)
    _reset_columns(doc)
    # the small map
    mx = M + left_w + 8
    mh = doc.h - y0 - M - 22
    _draw_map(doc, c["map"], mx, y0, map_w, mh, small=True)
    doc.set_xy(mx, y0 + mh + 1.5)
    _mini_legend(doc, mx, map_w)


def _mini_legend(doc, x, w):
    doc.set_font("Anek", "", SMALL)
    items = [("fill", STATUS["no_record"], "Not in register"), ("fill", STATUS["discrepancy"], "Differs"),
             ("line", PRIORITY_RGB["high"], "Possible dark stretch")]
    y = doc.get_y()
    cx = x
    for kind, col, lab in items:
        _swatch(doc, kind, col, cx, y + 0.6)
        doc.set_text_color(*INK2)
        doc.set_xy(cx + 7.5, y)
        tw = doc.get_string_width(lab) + 2
        doc.cell(tw, 5, lab)
        cx += 7.5 + tw + 3
    doc.set_xy(x, y + 5.5)
    doc.set_font("Anek", "", 8.5)
    doc.set_text_color(*INK3)
    doc.cell(w, 4, "Full map on page 3.")


def _swatch(doc, kind, col, x, y, level=None):
    if kind == "fill":
        doc.set_fill_color(*col)
        doc.set_draw_color(*[max(0, v - 40) for v in col])
        doc.set_line_width(0.15)
        doc.rect(x, y + 0.6, 5.5, 3.6, style="DF")
    else:
        lw = {"high": 2.4, "medium": 2.0, "low": 1.7}.get(level, 2.2)
        doc.set_draw_color(255, 255, 255)
        doc.set_line_width(lw + 0.7)
        doc.line(x, y + 2.4, x + 6, y + 2.4)
        doc.set_draw_color(*col)
        doc.set_line_width(lw)
        doc.line(x, y + 2.4, x + 6, y + 2.4)
        doc.set_draw_color(*DARK)
        doc.set_line_width(0.9)
        doc.line(x, y + 2.4, x + 6, y + 2.4)


def _bars(doc, x, y, w, title, sub, rows, colors=None, label_w=None, max_rows=10):
    """one labelled horizontal bar chart: label · bar · value. rows = [(label, value)]; one hue unless `colors`."""
    doc.set_xy(x, y)
    doc.set_font("Anek", "B", 12.5)
    doc.set_text_color(*INK)
    doc.cell(w, 6, title, new_x="LEFT", new_y="NEXT")
    if sub:
        doc.set_font("Anek", "", SMALL)
        doc.set_text_color(*INK3)
        doc.multi_cell(w, LH(SMALL), sub, new_x="LEFT", new_y="NEXT", align="L")
    doc.ln(1.5)
    if not rows:
        doc.set_font("Anek", "", BODY)
        doc.set_text_color(*INK3)
        doc.cell(w, 6, "Nothing to show.")
        return doc.get_y() + 6
    rows = rows[:max_rows]
    doc.set_font("Anek", "", 10)
    lw = label_w or min(w * 0.45, max(doc.get_string_width(str(r[0])) for r in rows) + 3)
    vmax = max(r[1] for r in rows) or 1
    val_w = 12
    bw = w - lw - val_w - 2
    yy = doc.get_y()
    bh, step = 5.0, 7.0
    for i, (lab, v) in enumerate(rows):
        doc.set_xy(x, yy)
        doc.set_text_color(*INK)
        lab = str(lab)
        while doc.get_string_width(lab) > lw - 2 and len(lab) > 4:
            lab = lab[:-2].rstrip() + "…" if not lab.endswith("…") else lab[:-2] + "…"
        doc.cell(lw, step - 1, lab)
        doc.set_fill_color(*((colors or {}).get(rows[i][0]) or BAR))
        L = max(0.8, bw * v / vmax) if v else 0
        if L:
            doc.rect(x + lw, yy + (step - 1 - bh) / 2, L, bh, style="F", round_corners=True, corner_radius=0.8)
        doc.set_xy(x + lw + L + 1.5, yy)
        doc.set_font("Anek", "B", 10)
        doc.cell(val_w, step - 1, fmt(v))
        doc.set_font("Anek", "", 10)
        yy += step
    return yy


def _all_numbers(doc, c, W):
    """every key number as the app shows it (the ribbon and More), one line above the footer of the charts page"""
    y = doc.h - M - 21
    doc.set_draw_color(*LINE)
    doc.set_line_width(0.2)
    doc.line(M, y - 1.5, M + W, y - 1.5)
    doc.set_xy(M, y)
    doc.set_font("Anek", "B", 10)
    doc.set_text_color(*INK)
    doc.cell(36, LH(10), "All key numbers:")
    doc.set_font("Anek", "", 10)
    doc.set_text_color(*INK2)
    doc.multi_cell(W - 36, LH(10), " · ".join(f"{n['label']} {fmt(n['value'])}" for n in c["numbers"]), align="L",
                   new_x="LMARGIN", new_y="NEXT")


def _page_charts(doc, c, W):
    doc.add_page()
    _all_numbers(doc, c, W)
    doc.set_xy(M, M)
    _kicker(doc, "Charts")
    _title(doc, "What the camera found, in numbers", 18, after=2)
    ch = c["charts"]
    gap = 10
    cw = (W - 2 * gap) / 3
    y = doc.get_y()
    match_col = {match_label("no_record"): STATUS["no_record"], match_label("discrepancy"): STATUS["discrepancy"],
                 match_label("matched"): STATUS["matched"], MATCHED_USE_UNKNOWN: STATUS["matched"]}
    k = c["kpis"]
    y1 = _bars(doc, M, y, cw, "Building use", f"{plural(k['buildings_analysed'], 'building')} · "
               f"use not known for {fmt(k['use_not_classified'])} (shown, never dropped)", ch["use"])
    y2 = _bars(doc, M + cw + gap, y, cw, "Floors counted",
               f"counted from the photo only: n = {fmt(ch['floors_n'])} ({fmt(ch['floors_left_out'])} estimates or not counted "
               "left out)", ch["floors"])
    y3 = _bars(doc, M + 2 * (cw + gap), y, cw, "Register match", "against the SYNTHETIC register", ch["match"], colors=match_col,
               label_w=min(cw * 0.6, 60))
    y = max(y1, y2, y3) + 8
    wide = 2 * cw + gap
    if c["street"]:
        st = c["tables"]["stretches"]["rows"]
        _bars(doc, M, y, wide, "Possible dark stretches on this street",
              f"recorded length in metres, in priority order; no lamp seen within 60 m (the detector finds about {c['recall']} "
              "of lamp heads)", [(f"#{r[0]} {r[1]} · {r[4]}", int(str(r[5]).replace(',', ''))) for r in st],
              colors={f"#{r[0]} {r[1]} · {r[4]}": PRIORITY_RGB.get(r[1].lower(), BAR) for r in st}, label_w=min(wide * 0.45, 90))
        _shops_chart(doc, c, M + wide + gap, y, cw)
        _reset_columns(doc)
        return
    # findings by street: two series, side by side per street (not stacked), with a legend
    doc.set_xy(M, y)
    doc.set_font("Anek", "B", 12.5)
    doc.set_text_color(*INK)
    doc.cell(wide, 6, "Findings by street", new_x="LEFT", new_y="NEXT")
    doc.set_font("Anek", "", SMALL)
    doc.set_text_color(*INK3)
    doc.cell(wide, LH(SMALL), "buildings not in the register and buildings that differ, per street (synthetic register)",
             new_x="LEFT", new_y="NEXT")
    ly = doc.get_y() + 1
    lx = M
    for col, lab in ((STATUS["no_record"], "Not in the register"), (STATUS["discrepancy"], "Differs from the register")):
        _swatch(doc, "fill", col, lx, ly)
        doc.set_xy(lx + 7, ly)
        doc.set_font("Anek", "", 10)
        doc.set_text_color(*INK2)
        doc.cell(doc.get_string_width(lab) + 2, 5, lab)
        lx += 7 + doc.get_string_width(lab) + 8
    yy = ly + 7
    fit = max(1, int((doc.h - M - 30 - yy) // 9))           # rows that fit above the key-numbers line
    rows = [r for r in ch["by_street"] if r[1] + r[2]][:fit]
    if not rows:
        doc.set_xy(M, yy)
        doc.set_font("Anek", "", BODY)
        doc.cell(wide, 6, "No building here is missing from or differs from the register.")
    else:
        doc.set_font("Anek", "", 10)
        lw = min(wide * 0.38, max(doc.get_string_width(r[0]) for r in rows) + 3)
        vmax = max(max(r[1], r[2]) for r in rows) or 1
        bw = (wide - lw - 14)
        for st, a, d, _ in rows:
            doc.set_xy(M, yy)
            doc.set_text_color(*INK)
            doc.cell(lw, 8, st)
            for j, (v, col) in enumerate(((a, STATUS["no_record"]), (d, STATUS["discrepancy"]))):
                L = bw * v / vmax
                by_ = yy + 0.6 + j * 3.6
                if v:
                    doc.set_fill_color(*col)
                    doc.rect(M + lw, by_, max(0.8, L), 3.0, style="F")
                doc.set_xy(M + lw + (L if v else 0) + 1.2, by_ - 0.9)
                doc.set_font("Anek", "B", 9)
                doc.cell(10, 4.6, fmt(v))
                doc.set_font("Anek", "", 10)
            yy += 9
        more = len([r for r in ch["by_street"] if r[1] + r[2]]) - len(rows)
        if more > 0:
            doc.set_xy(M, yy)
            doc.set_font("Anek", "", SMALL)
            doc.set_text_color(*INK3)
            doc.cell(wide, 5, f"+ {plural(more, 'more street')} with findings (full list in the Excel file)")
    _shops_chart(doc, c, M + wide + gap, y, cw)
    _reset_columns(doc)


def _shops_chart(doc, c, x, y, cw):
    ch = c["charts"]
    if ch["shops"]:
        sh = c["shops"]
        yb = _bars(doc, x, y, cw, "Businesses vs OpenStreetMap",
                   f"our camera found {fmt(sh['counts']['camera'])}; OpenStreetMap lists {fmt(sh['counts']['osm'])} along "
                   f"these streets (≤ {sh['scope_m']} m)", ch["shops"], label_w=46)
        doc.set_xy(x, yb + 2)
        _text(doc, f"Same place = within {sh['match_m']} m; {fmt(sh['counts']['matched_same_name'])} of the "
                   f"{fmt(sh['counts']['matched'])} matched also have the same name. OpenStreetMap is crowd-sourced, not an "
                   "official register.", size=SMALL, color=INK3, w=cw)
    else:
        doc.set_xy(x, y)
        _text(doc, "Businesses vs OpenStreetMap: " + c["shops"].get("note", "not available"), size=SMALL, color=INK3, w=cw)


def _page_map(doc, c, W, actions=False):
    doc.add_page()
    _kicker(doc, "Map · what to do next" if actions else "Map")
    _title(doc, "Where the findings are, and what to do" if actions else "Where the findings are", 18, after=1)
    y = doc.get_y()
    leg_w = 66 if not actions else 132
    mw = W - leg_w - 8
    mh = doc.h - y - M - 8
    _draw_map(doc, c["map"], M, y, mw, mh)
    x = M + mw + 8
    doc.set_xy(x, y)
    _in_column(doc, x, leg_w)
    doc.set_font("Anek", "B", 12.5)
    doc.set_text_color(*INK)
    doc.cell(leg_w, 6, "Key", new_x="LMARGIN", new_y="NEXT")
    doc.ln(1)
    items = [("fill", STATUS["no_record"], "Building not in the register", None),
             ("fill", STATUS["discrepancy"], "Building differs from the register", None),
             ("fill", STATUS["matched"], "Building matches / entry exists", None),
             ("line", PRIORITY_RGB["high"], "Possible dark stretch, High priority", "high"),
             ("line", PRIORITY_RGB["medium"], "… Medium priority", "medium"),
             ("line", PRIORITY_RGB["low"], "… Low priority", "low")]
    for kind, col, lab, lv in items:
        yy = doc.get_y()
        _swatch(doc, kind, col, x, yy + 0.4, lv)
        doc.set_xy(x + 8, yy)
        doc.set_font("Anek", "", 10.5)
        doc.set_text_color(*INK2)
        doc.multi_cell(leg_w - 8, LH(10.5), lab, new_x="LMARGIN", new_y="NEXT", align="L")
        doc.ln(1.2)
    doc.ln(2)
    _text(doc, "The numbers on the dark stretches are their place in the priority list (appendix). The wider and deeper "
               "the edge, the higher the priority.", size=SMALL, color=INK2, w=leg_w)
    _text(doc, f"Roads: {c['map']['road_source']}. Building outlines: OpenStreetMap (as analysed). Drawn from the app's "
               "own data: no Google map or photo is reproduced.", size=SMALL, color=INK3, w=leg_w)
    _text(doc, "© OpenStreetMap contributors", size=SMALL, color=INK3, w=leg_w)
    if actions:
        doc.ln(2)
        doc.set_font("Anek", "B", 13)
        doc.set_text_color(*INK)
        doc.cell(leg_w, 6, "What to do next", new_x="LMARGIN", new_y="NEXT")
        doc.ln(1)
        for i, a in enumerate(c["actions"], 1):
            doc.set_font("Anek", "B", BODY)
            doc.set_text_color(*INK)
            doc.multi_cell(leg_w, LH(BODY), f"{i}. {a['what']}", new_x="LMARGIN", new_y="NEXT", align="L")
            doc.set_font("Anek", "", 10)
            doc.set_text_color(*INK2)
            doc.multi_cell(leg_w, LH(10), a["why"], new_x="LMARGIN", new_y="NEXT", align="L")
            doc.ln(1.6)
    _reset_columns(doc)


def _page_actions(doc, c, W):
    doc.add_page()
    _kicker(doc, "What to do next")
    _title(doc, "The top actions, most important first", 18, after=3)
    if not c["actions"]:
        _text(doc, "Nothing stands out here: no possible dark stretch, no register finding and no item waiting for review.")
        return
    for i, a in enumerate(c["actions"], 1):
        y = doc.get_y()
        doc.set_fill_color(*SODIUM)
        doc.circle(x=M + 3.6, y=y + 3.4, radius=3.4, style="F")
        doc.set_xy(M, y + 0.6)
        doc.set_font("Anek", "B", 10.5)
        doc.set_text_color(255, 255, 255)
        doc.cell(7.2, 5.6, str(i), align="C")
        doc.set_xy(M + 11, y)
        doc.set_font("Anek", "B", 12.5)
        doc.set_text_color(*INK)
        doc.multi_cell(W - 11, LH(12.5), a["what"], new_x="LEFT", new_y="NEXT", align="L")
        doc.set_font("Anek", "", BODY)
        doc.set_text_color(*INK2)
        doc.multi_cell(W - 11, LH(BODY), a["why"], new_x="LMARGIN", new_y="NEXT", align="L")
        doc.ln(2.2)
    _text(doc, "Every action lists its street; the appendix and the Excel file list each building, pole and stretch with a "
               "Google Maps link.", size=SMALL, color=INK3)


def _section(doc, title):
    doc.ln(1)
    doc.set_font("Anek", "B", 12.5)
    doc.set_text_color(*INK)
    doc.multi_cell(0, 6, title, new_x="LMARGIN", new_y="NEXT", align="L")
    doc.ln(0.5)


def _page_method(doc, c, W):
    doc.add_page()
    _kicker(doc, "Method & limits")
    _title(doc, "How this was made, and what it can't tell you", 18, after=2)
    gap = 10
    cw = (W - gap) / 2
    y0 = doc.get_y()
    # left column: how it works, data sources, cost
    _in_column(doc, M, cw)
    doc.set_y(y0)
    _section(doc, "How it works")
    _text(doc, "Google Street View photos of every analysed street are read by small models on our own computer (they find "
               "buildings, poles, streetlights and shop signs, and read the signs). Only unsure cases go to a cloud AI model. "
               "Each finding is placed on the map and compared with the property and asset register.", size=10.5)
    _section(doc, "Data sources")
    for k_, v in c["sources"]:
        _text(doc, f"{k_}: {v}", size=10, after=0.8)
    if c["cost"] and c["cost"]["line"]:
        _section(doc, "Cost of this run" + (" (whole area)" if c["street"] else ""))
        _text(doc, c["cost"]["line"], size=10)
    yl = doc.get_y()
    # right column: limits, Gate 1, projection
    x2 = M + cw + gap
    _in_column(doc, x2, cw)
    doc.set_y(y0)
    _section(doc, "Limits")
    for t in c["limits"]:
        _text(doc, "• " + t, size=10, after=0.8)
    _section(doc, c["gate1"]["title"])
    _text(doc, c["gate1"]["status"] + ". " + c["gate1"]["note"], size=10, color=INK, after=0.8)
    _text(doc, c["gate1"]["row"], size=10)
    _reset_columns(doc)
    _page_scale(doc, c, W)


def _are(n):
    return "is" if n == 1 else "are"


def _page_scale(doc, c, W):
    """page 6: the city-scale projection, then floor-count confidence and the OpenStreetMap cross-checks"""
    doc.add_page()
    _kicker(doc, "Method & limits · scale and confidence")
    _title(doc, "What a whole city would take, and how sure the counts are", 18, after=2)
    _projection_block(doc, c, W)
    gap = 10
    cw = (W - gap) / 2
    y0 = doc.get_y() + 2
    _in_column(doc, M, cw)
    doc.set_y(y0)
    _section(doc, "Floor-count confidence")
    _text(doc, c["floor_rule"] + (f" The AI floor count was {c['floor_check']}." if c.get("floor_check") else ""), size=10)
    lv = c["levels"]
    if lv.get("available"):
        _text(doc, f"OpenStreetMap cross-check{' (whole area)' if c['street'] else ''}: {fmt(lv['tagged'])} of the {plural(lv['buildings'], 'analysed building')} carry "
                   f"a building:levels tag; {('our count agrees exactly for ' + fmt(lv['exact']) + ' and within one floor for ' + fmt(lv['within_1']) + ' of the ' + fmt(lv['compared']) + ' we counted') if lv['compared'] else 'none of them has a floor count from us'}. "
                   "Too few to judge accuracy; our counts are never changed by it.", size=10)
    x2 = M + cw + gap
    _in_column(doc, x2, cw)
    doc.set_y(y0)
    _section(doc, "Businesses vs OpenStreetMap")
    sh = c["shops"]
    if sh.get("available"):
        n = sh["counts"]
        _text(doc, f"Our camera found {plural(n['camera'], 'shop or business', 'shops and businesses')}; OpenStreetMap lists "
                   f"{fmt(n['osm'])} along {'this street' if c['street'] else 'these streets'}. {fmt(n['matched'])} "
                   f"{_are(n['matched'])} the same place (within {sh['match_m']} m; {fmt(n['matched_same_name'])} also with the same "
                   f"name), {fmt(n['camera_only'])} {_are(n['camera_only'])} seen by our camera only and {fmt(n['osm_only'])} "
                   f"{_are(n['osm_only'])} in OpenStreetMap only.", size=10)
        _text(doc, sh["rule"], size=9.5, color=INK3)
        _text(doc, sh["note"], size=9.5, color=INK3)
    else:
        _text(doc, sh.get("note", "Not available."), size=10)
    _reset_columns(doc)


def _projection_block(doc, c, W):
    p = c["projection"]
    _section(doc, "City-scale projection — an estimate, as a range from our runs")
    if not p.get("available"):
        _text(doc, "Not available: " + p.get("note", ""), size=10)
        return
    rows = []
    for r in p["cities"]:
        lo, hi = r["photos"], r["usd"]
        h = r["gpu_hours"]
        rows.append([("» " if c.get("city") == r["name"] else "") + f"Whole {r['name']}", f"about {fmt(r['km'])} km",
                     _range(lo["low"], lo["high"], million_from=1e5), _range(hi["low"], hi["high"], "$"),
                     f"{fmt(h['low'])} – {fmt(h['high'])} h" if h else "—",
                     f"{fmt(r['colab_days']['low'])} – {fmt(r['colab_days']['high'])}" if r.get("colab_days") else "—"])
    _table(doc, {"columns": ["City (local map data)", "Streets", "Street View photos", "Cost (list price)", "GPU time",
                             "Colab days (2.5 h/day)"], "rows": rows, "links": [None] * len(rows)},
           (62, 34, 46, 46, 40, 41), size=TABLE, link_col=False)
    for a in p["assumptions"]:
        _text(doc, "• " + a, size=9.5, color=INK2, after=0.4)


def _sig2(v):
    """two significant figures: 4,426 → 4,400; 18,265 → 18,000"""
    if v <= 0:
        return 0
    d = int(math.floor(math.log10(v))) - 1
    return round(v, -d) if d > 0 else round(v)


def _range(lo, hi, unit="", million_from=1e6):
    """a projection range in one unit, two significant figures: 0.63 – 2.6 million · $4,400 – $18,000"""
    if hi >= million_from:
        return f"{unit}{_sig2(lo) / 1e6:g} – {unit}{_sig2(hi) / 1e6:g} million"
    return f"{unit}{fmt(_sig2(lo))} – {unit}{fmt(_sig2(hi))}"


def _k(v):
    return f"{_sig2(v) / 1e6:g} million" if v >= 1e6 else fmt(_sig2(v))


APPENDIX = {  # key columns per table (indexes into the full row); the Excel file keeps every column
    "stretches": ([0, 1, 2, 4, 5, 3], (10, 22, 110, 62, 26, 39)),
    "buildings": ([0, 1, 3, 4, 5, 8, 10, 12], (26, 50, 26, 26, 24, 42, 44, 31)),
    "assets": ([0, 1, 2, 4, 5], (28, 36, 62, 70, 73)),
    "review": ([0, 1, 2, 4, 5], (22, 60, 58, 98, 31)),
    "shops": ([0, 1, 2, 5, 6, 7], (64, 30, 60, 50, 43, 22)),
}


def _appendix(doc, c, W):
    T = c["tables"]
    doc.add_page()
    _kicker(doc, "Appendix")
    _title(doc, "The lists behind the numbers", 18, after=0.5)
    _text(doc, "Key columns only. Every column, with a Google Maps link on each row, is in the Excel file of this report.",
          size=10.5, color=INK2, after=2)
    first = True
    for key, title in (("stretches", T["stretches"]["title"]), ("buildings", T["buildings"]["title"]),
                       ("assets", T["assets"]["title"]), ("review", T["review"]["title"]),
                       ("shops", T["shops"]["title"])):
        t = T[key]
        idx, widths = APPENDIX[key]
        if not first and doc.get_y() > doc.h - 60:
            doc.add_page()
        first = False
        doc.set_font("Anek", "B", 12.5)
        doc.set_text_color(*INK)
        doc.cell(0, 7, f"{title} ({fmt(len(t['rows']))})", new_x="LMARGIN", new_y="NEXT")
        if key == "shops" and c["shops"].get("available"):
            _text(doc, c["shops"]["note"], size=9.5, color=INK3, after=0.5)
        rows, rest = t["rows"], None
        if key == "review":
            rows = [r for r in t["rows"] if (r[0] or 9) <= 5]
            if len(rows) < len(t["rows"]):
                rest = (f"{plural(len(t['rows']) - len(rows), 'more item')} of priority 6 (poles and streetlights seen in one "
                        "photo only, a second look) are listed in the Excel file (sheet Review items).")
        if key == "shops":
            rows = [r for r in t["rows"] if not r[0].startswith("Seen by our camera")]
            if len(rows) < len(t["rows"]):
                rest = (f"{plural(len(t['rows']) - len(rows), 'business', 'businesses')} seen by our camera but not in "
                        "OpenStreetMap are listed in the Excel file (sheet OSM shops).")
        short = lambda col, v: v.split(":")[0] if col == "Floor-count confidence" and isinstance(v, str) else v
        sub = {"columns": [t["columns"][i] for i in idx], "rows": [[short(t["columns"][i], r[i]) for i in idx] for r in rows],
               "links": [None] * len(rows)}
        _table(doc, sub, widths, size=TABLE, link_col=False)
        if rest:
            doc.ln(1)
            _text(doc, rest, size=10, color=INK2)
        doc.ln(3)


# ---------------------------------------------------------------------------------------------------------- Excel
def xlsx(c):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "About"
    bold = Font(bold=True)
    ws.append(["GEO-CASCADIA report", c["title"]])
    ws.append(["Scope", c["scope"]])
    ws.append(["Made", c["generated"]])
    ws.append(["Register", "SYNTHETIC demo data (made up), not the city's records"])
    ws.append([])
    ws.append(["Key number", "Value", "Note"])
    for n in c["numbers"]:
        ws.append([n["label"], n["value"]] + ([n["note"]] if n.get("note") else []))
    ws.append([])
    ws.append(["Data source", "Detail"])
    for k, v in c["sources"]:
        ws.append([k, v])
    ws.append([])
    ws.append(["OpenStreetMap cross-checks", ""])
    sh, lv = c["shops"], c["levels"]
    if sh.get("available"):
        n = sh["counts"]
        ws.append(["Businesses found by our camera", n["camera"]])
        ws.append(["OpenStreetMap businesses along these streets", n["osm"]])
        ws.append(["Matched (same place)", n["matched"], f"{n['matched_same_name']} also with the same name"])
        ws.append(["Seen by our camera, not in OpenStreetMap", n["camera_only"]])
        ws.append(["In OpenStreetMap, not seen by our camera", n["osm_only"]])
        ws.append(["How they are compared", sh["rule"]])
        ws.append(["Note", sh["note"]])
    else:
        ws.append(["Businesses vs OpenStreetMap", sh.get("note")])
    if lv.get("available"):
        ws.append(["Analysed buildings with OSM building:levels", lv["tagged"], f"of {lv['buildings']}"])
        ws.append(["... of those we counted: same floors / within one", f"{lv['exact']} / {lv['within_1']}", f"of {lv['compared']}"])
        ws.append(["Note", lv["note"]])
    ws.append(["Floor-count confidence", c["floor_rule"]])
    ws.append([])
    p = c["projection"]
    ws.append(["City-scale projection (an ESTIMATE)", "low – high, from our completed runs"])
    if p.get("available"):
        for r in p["cities"]:
            h = r["gpu_hours"]
            ws.append([f"Whole {r['name']}", f"about {fmt(r['km'])} km of streets → {_k(r['photos']['low'])} – {_k(r['photos']['high'])} "
                                             f"photos, ${_k(r['usd']['low'])} – ${_k(r['usd']['high'])}"
                                             + (f", {fmt(h['low'])} – {fmt(h['high'])} GPU hours" if h else "")])
        for a in p["assumptions"]:
            ws.append(["", a])
    else:
        ws.append(["", p.get("note")])
    ws.append([])
    ws.append(["Limits", ""])
    for t in c["limits"]:
        ws.append(["", t])
    ws.append([])
    ws.append(["Footer", FOOTER])
    for row in ws.iter_rows():
        if row[0].value in ("GEO-CASCADIA report", "Key number", "Data source", "Limits", "OpenStreetMap cross-checks",
                            "City-scale projection (an ESTIMATE)"):
            for cell in row:
                cell.font = bold
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 110
    head_fill = PatternFill("solid", fgColor="ECE8DE")
    for key in ("buildings", "assets_all", "stretches", "review", "shops"):
        t = c["tables"][key]
        sh = wb.create_sheet(t["sheet"])
        sh.append(t["columns"])
        for cell in sh[1]:
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row, link in zip(t["rows"], t["links"]):
            sh.append(row)
            if link:
                cell = sh.cell(row=sh.max_row, column=len(row))
                cell.hyperlink = link
                cell.font = Font(color="9C4E07", underline="single")
        for i, col in enumerate(t["columns"], 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in t["rows"]]) + 2
            sh.column_dimensions[get_column_letter(i)].width = min(60, max(8, width))
        sh.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filename(c, ext):
    import re
    base = c["slug"] + (("_" + re.sub(r"[^a-z0-9]+", "_", c["street"].lower()).strip("_")) if c["street"] else "")
    return f"geo-cascadia_report_{base}_{c['generated']}.{ext}"
