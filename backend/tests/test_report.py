"""D54 lighting priority + D55 area report: every number in the report equals what the app shows (the API the app reads),
for Ward 29 and one small area (Vadakku Masi Veethi, a live run in Madurai), for the whole area and for one street.
The PDF's text and the Excel sheets are read back and compared with the same values."""
import io

import pytest

from app import lighting, mapdata, report

AREAS = ["ward29", "vadakku_masi_veethi_f17937"]
STREET = {"ward29": "Sathy Main Road", "vadakku_masi_veethi_f17937": None}


@pytest.fixture(scope="module")
def api(online):
    """the online app with the map data pointed at its own database (the session's offline app re-points the module)"""
    mapdata.configure(online.app.state.data.pool)
    mapdata._DOWN["until"] = 0
    return online


_BUNDLES = {}


def _content(api, slug, street=None):
    app = api.app
    mapdata.configure(app.state.data.pool)
    mapdata._DOWN["until"] = 0
    s = app.state.data.db
    _BUNDLES[slug] = s.bundle(slug)
    return report.content(s, _BUNDLES[slug], app.state.runfiles.get(slug), app.state.model_card.get(),
                          app.state.settings.areas_dir, street=street)


# ------------------------------------------------------------------------------------------------ the priority rule
def test_rule_points_and_levels():
    assert [lighting.length_points(m) for m in (60, 119, 120, 239, 240, 900)] == [1, 1, 2, 2, 3, 3]
    assert [lighting.activity_points(n) for n in (0, 1, 4, 5, 9, 10)] == [0, 1, 1, 2, 2, 3]
    assert [lighting.road_kind(h) for h in ("trunk", "secondary_link", "tertiary", "unclassified", "residential",
                                            "service", None)] == ["main", "main", "connecting", "connecting",
                                                                  "residential", "residential", "unknown"]
    assert [lighting.level(s) for s in (9, 7, 6, 5, 4, 0)] == ["high", "high", "medium", "medium", "low", "low"]
    r = lighting.score_row("g", 140, "primary", 5, 3, 1)
    assert r["reason"] == "140 m on a main road, 9 shops and businesses along it" and r["score"] == 2 + 3 + 2
    assert lighting.score_row("g", 70, None, 0, 0, 0)["reason"] == "70 m on a road of unknown type, no shops or businesses along it"
    assert lighting.score_row("g", 70, "service", 1, 0, 0)["reason"].endswith("1 shop or business along it")


def test_ward29_priority_table(api):
    rows = api.get("/areas/ward29/lighting").json()
    assert rows["available"] and len(rows["rows"]) == 11
    for r in rows["rows"]:
        assert r["score"] == sum(r["points"].values())
        assert r["level"] == lighting.level(r["score"])
        assert r["activity"]["total"] == r["activity"]["buildings"] + r["activity"]["osm_points"] + r["activity"]["sign_businesses"]
    order = [(r["score"], r["length_m"]) for r in rows["rows"]]
    assert order == sorted(order, key=lambda x: (-x[0], -x[1]))
    top = rows["rows"][0]
    assert top["id"] == "gap60-001" and top["road_class"] == "trunk" and top["level"] == "high"
    # the map features and a typed question carry the same priority as the table
    feats = {f["properties"]["id"]: f["properties"] for f in api.get("/areas/ward29/geojson?layers=gaps").json()["features"]}
    for r in rows["rows"]:
        assert (feats[r["id"]]["priority"], feats[r["id"]]["priority_reason"]) == (r["level"], r["reason"])
    q = api.post("/query", json={"area": "ward29", "text": "Show high priority dark stretches"}).json()
    assert q["understanding"]["status"] == "ok" and q["parsed_filters"]["priority"] == "high"
    assert {x["id"] for x in q["rows"]} == {r["id"] for r in rows["rows"] if r["level"] == "high"}
    chips = api.post("/query", json={"area": "ward29", "filters": {"intent": "streetlight_gaps", "interval_m": 60, "priority": "low"}}).json()
    assert chips["total"] == sum(r["level"] == "low" for r in rows["rows"])
    # a building question does not take a priority level: it is reported as ignored, never applied silently
    b = api.post("/query", json={"area": "ward29", "text": "high priority commercial buildings"}).json()
    assert b["understanding"]["status"] == "partial" and "high priority" in b["understanding"]["ignored"]
    # "possible dark stretches" is the app's own wording: nothing ignored
    p = api.post("/query", json={"area": "ward29", "text": "Show possible dark stretches"}).json()
    assert p["understanding"]["status"] == "ok" and p["total"] == 11


def test_priority_offline_is_said_plainly(offline):
    r = offline.get("/areas/ward29/lighting").json()
    assert r["available"] is False and r["note"] == lighting.OFFLINE_NOTE
    q = offline.post("/query", json={"area": "ward29", "text": "Show high priority dark stretches"}).json()
    assert q["total"] is None and q["rows"] == [] and "not available" in q["why_empty"][-1]["step"]


# --------------------------------------------------------------------------------------- report numbers = the app's
@pytest.mark.parametrize("slug", AREAS)
def test_report_numbers_equal_the_app(api, slug):
    c = _content(api, slug)
    area = api.get(f"/areas/{slug}").json()
    kpi = area["dashboard"]["kpi"]
    k = c["kpis"]
    for key in ("buildings_analysed", "unmatched_properties", "buildings_with_discrepancy", "streetlights", "poles",
                "named_businesses", "sign_text_unverified", "names_confirmed_by_google", "low_confidence_observations",
                "unmapped_businesses", "use_not_classified", "waiting_for_review", "streets_covered"):
        assert k[key] == kpi[key], key
    gaps = api.get(f"/areas/{slug}/geojson?layers=gaps").json()["features"]      # the app counts the map's stretches
    assert k["streetlight_gaps"] == len(gaps)
    assert [n["value"] for n in c["numbers"]] == [k[d[0]] for d in report.KPI_DEFS]

    T = c["tables"]
    assert len(T["buildings"]["rows"]) == kpi["unmatched_properties"] + kpi["buildings_with_discrepancy"]
    review = api.get(f"/review?area={slug}").json()
    items = list(range(review["total"]))
    assert len(T["review"]["rows"]) == len(items) == kpi["low_confidence_observations"]
    geo = api.get(f"/areas/{slug}/geojson?layers=assets").json()["features"]
    assert len(T["assets"]["rows"]) == sum(f["properties"]["register_status"] in ("unrecorded_asset", "discrepancy") for f in geo)
    # D56: the Excel sheet lists every pole and streetlight; the finding rows come first and equal the PDF's
    assert len(T["assets_all"]["rows"]) == len(geo) == kpi["streetlights"] + kpi["poles"]
    assert T["assets_all"]["rows"][:len(T["assets"]["rows"])] == T["assets"]["rows"]
    assert all(r[5] == "" for r in T["assets_all"]["rows"][len(T["assets"]["rows"]):])

    light = api.get(f"/areas/{slug}/lighting").json()["rows"]
    assert [(r[3], r[1], r[2]) for r in T["stretches"]["rows"]] == \
        [(x["id"], report.PRIORITY_WORD[x["level"]], x["reason"]) for x in light]

    hood = api.get(f"/areas/{slug}/hood").json()
    rt = hood["routing"]
    sv, ai = rt["street_view"]["usd"], rt["totals"]["usd"]
    # D60: Google and AWS on separate lines, no combined total; the India billing line under the Google one
    L = c["cost"]["lines"]
    assert L[0].startswith("Street View photos (Google): " + report.usd_text(sv)) and report.fmt(rt["street_view"]["photos"]) in L[0]
    assert L[1] == hood["billing"]["line"] and L[2].startswith("Cloud AI (Amazon Nova Lite, AWS): " + report.usd_text(ai))
    assert report.usd_text(sv + ai) not in c["cost"]["line"]
    im = hood["imagery"]
    assert c["sources"][0][1].startswith(f"taken {report.month_text(im['oldest'])} – {report.month_text(im['newest'])}; "
                                         f"{im['older_than_cutoff']} of {im['camera_stops_dated']}")

    g = api.get("/model-card").json().get("gate1_position") or api.app.state.model_card.get()["gate1_position"]
    assert c["gate1"]["status"] == "Status: Not verified" and c["gate1"]["note"].startswith(g["status_note"])
    fair = next((v for kk, v in (g["vs OSM front-wall centre"].get(slug) or {}).items() if kk.startswith("camera-derived")), None)
    if fair:
        assert f"n = {fair['n']}, median {fair['median_m']} m" in c["gate1"]["row"]
    else:
        assert "not in Trust's Gate 1 table" in c["gate1"]["row"]

    _files_carry_the_same_values(c)


@pytest.mark.parametrize("slug", ["ward29"])
def test_street_report_equals_the_app(api, slug):
    street = STREET[slug]
    c = _content(api, slug, street)
    k = c["kpis"]

    def total(f):
        return api.post("/query", json={"area": slug, "filters": {**f, "street": street}}).json()["total"]
    assert k["buildings_analysed"] == total({"intent": "buildings"})
    assert k["unmatched_properties"] == total({"intent": "buildings", "match_status": "no_record"})
    assert k["buildings_with_discrepancy"] == total({"intent": "buildings", "match_status": "discrepancy"})
    assert k["streetlight_gaps"] == total({"intent": "streetlight_gaps", "interval_m": 60})
    assert k["streetlights"] == total({"intent": "assets", "asset_type": "streetlight"})
    assert k["poles"] == total({"intent": "assets", "asset_type": "pole"})
    assert all(r[2] == street for r in c["tables"]["review"]["rows"])
    assert all(r[1] == street for r in c["tables"]["buildings"]["rows"])
    _files_carry_the_same_values(c)
    assert api.get(f"/areas/{slug}/report.pdf?street=Nowhere Street").status_code == 404


def _files_carry_the_same_values(c):
    import pypdfium2
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(report.xlsx(c)))
    for key in ("buildings", "assets_all", "stretches", "review", "shops"):
        t = c["tables"][key]
        sh = wb[t["sheet"]]
        got = [[cell.value for cell in row] for row in sh.iter_rows(min_row=1)]
        assert got[0] == t["columns"]
        assert got[1:] == [[v if v != "" else None for v in r] for r in t["rows"]]     # Excel keeps a blank cell empty
        for i, link in enumerate(t["links"], start=2):
            assert sh.cell(row=i, column=len(t["columns"])).hyperlink.target == link
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows() if r[0].value}
    for n in c["numbers"]:
        assert about[n["label"]] == n["value"]
    doc = pypdfium2.PdfDocument(report.pdf(c))
    text = "".join(doc[i].get_textpage().get_text_range() for i in range(len(doc))).replace("\r\n", " ").replace("\n", " ")
    flat = " ".join(text.split())
    for n in c["numbers"]:
        assert f"{n['label']} {report.fmt(n['value'])}" in flat or report.fmt(n["value"]) in flat
    for r in c["tables"]["stretches"]["rows"]:
        assert r[3] in flat                                             # every stretch id is in the PDF
    for r in c["tables"]["buildings"]["rows"]:
        assert r[0] in flat
    assert "SYNTHETIC" in flat and "© OpenStreetMap contributors" in flat
    _gis_carries_the_same_values(c)


def _gis_carries_the_same_values(c):
    """extras 1: the GeoJSON and the zipped Shapefile carry the Excel sheets' rows (same count per layer, same values;
    Shapefile names via fields.csv), in WGS84 with a .prj"""
    import csv
    import json
    import tempfile
    import zipfile
    import shapefile
    from app import gisexport
    b = _BUNDLES[c["slug"]]
    gj = json.loads(gisexport.geojson(c, b))
    z = zipfile.ZipFile(io.BytesIO(gisexport.shapefile_zip(c, b)))
    key = {}
    for row in list(csv.reader(io.StringIO(z.read("fields.csv").decode("utf-8-sig"))))[1:]:
        key.setdefault(row[0], {})[row[1]] = row[2]
    d = tempfile.mkdtemp()
    z.extractall(d)
    for layer, tkey, _ in gisexport.LAYERS:
        t = c["tables"][tkey]
        feats = [f for f in gj["features"] if f["properties"]["layer"] == layer]
        assert len(feats) == len(t["rows"]), layer
        for f, row, link in zip(feats, t["rows"], t["links"]):
            want = dict(zip(t["columns"], row))
            want[t["columns"][-1]] = link or want[t["columns"][-1]]
            assert {k: v for k, v in f["properties"].items() if k != "layer"} == want
            assert f["geometry"]["type"] in ("Point", "Polygon", "LineString")
        assert z.read(f"{layer}.prj").decode().startswith('GEOGCS["GCS_WGS_1984"')
        r = shapefile.Reader(f"{d}/{layer}")
        names = [fl.name for fl in r.fields[1:]]
        assert [key[layer][n] for n in names] == t["columns"]
        recs = r.records()
        assert len(recs) == len(t["rows"]), layer
        for rec, row, link in zip(recs, t["rows"], t["links"]):
            for col, v in zip(t["columns"], list(rec)):
                exp = link if col == t["columns"][-1] and link else dict(zip(t["columns"], row))[col]
                if exp is None or exp == "":
                    assert v in (None, ""), (layer, col, v)
                elif isinstance(exp, float):
                    assert abs(float(v) - exp) < 0.01
                else:
                    assert (str(v) == str(exp)) or (isinstance(exp, str) and len(exp.encode()) > 254 and exp.startswith(v)), (layer, col, v, exp)
        r.close()


def test_report_offline_builds_and_says_so(offline):
    """D4 offline data mode: the report still builds from the JSON copy; priority is said to be unavailable"""
    r = offline.get("/areas/ward29/report.xlsx")
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb["Possible dark stretches"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 11 and all(x[1] == "not available" and x[2] == lighting.OFFLINE_NOTE for x in rows)
    assert len(list(wb["Assets"].iter_rows(min_row=2))) == 262                 # D64 re-run
    assert offline.get("/areas/ward29/report.pdf?street=Sathy Main Road").status_code == 200


def test_camera_only_count_is_shown_next_to_the_building_count(api):
    """D56: "+ N seen only by camera" comes from building_positions.json no_footprint, is the same in the area card, Hood
    and the report, and is never added to the building count"""
    cards = {a["slug"]: a for a in api.get("/areas").json()["areas"]}
    expect = {"ward29": 47, "trichy_bharathidasan_salai": 79, "tiruppur_uthukuli_road": 12}    # ward29: D64 re-run (Sep run: 9)
    for slug, n in expect.items():
        assert cards[slug]["counts"]["camera_only_buildings"] == n
        assert api.get(f"/areas/{slug}/camera-buildings").json()["count"] == n
        assert api.get(f"/areas/{slug}/hood").json()["n"]["camera_only_buildings"] == n
    c = _content(api, "ward29")
    assert c["numbers"][0]["value"] == 373 and c["numbers"][0]["note"] == "+ 47 seen only by camera (no map outline)"   # D64 re-run
    assert _content(api, "ward29", "Sathy Main Road")["numbers"][0].get("note") == ""      # no street on camera-only points
